from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
sections = [
('Project overview and objectives', '''This project will build a dashboard for one individual using a React frontend, a Python backend and a database. A Fitbit Inspire 3 records sleep, heart rate and steps. The participant synchronizes the band with the Google Health application through their Google account. The backend obtains authorized JSON data, stores it, calculates summaries and serves daily and weekly views to the dashboard.

The observation period is 14 consecutive local calendar days. The dashboard should help the participant and project team describe sleep patterns, daily activity and heart rate changes during that period. It is a descriptive student project; two weeks from one person cannot establish population trends or prove that activity causes better sleep.

Recommended implementation: React and TypeScript with Vite; FastAPI and Python; PostgreSQL accessed through SQLAlchemy; a scheduled Python ingestion command; and Recharts for presentation. Keep data collection separate from dashboard requests so the dashboard remains usable when Google is unavailable.

Primary objectives
- Collect authorized sleep, heart rate and step data for the configured participant and observation window.
- Preserve source records and normalized records in a database with repeatable, duplicate safe ingestion.
- Display daily details, two study weeks and a complete 14 day overview.
- Explain missing values, data coverage and the last successful synchronization.
- Produce reproducible calculations and a final study summary with clear limitations.

Scope and deliverables
The minimum product includes account connection or authorized JSON import, daily and weekly screens, persistent storage, synchronization status, date filtering and an export. Deliver the source repository, database migrations, setup guide, configuration example, tests and project report. Optional extensions include intraday step charts, a sleep stage timeline and exploratory comparisons. Machine learning, medical alerts, multiple participants and continuous long term monitoring are outside the initial scope.

This document is a proposed implementation specification, dated 7 October 2026. No participant measurements, API credentials or access permissions have been verified for this project.'''),
('Data access and feasibility', '''Confirm the data integration before beginning the 14 day recording period. Having a Google account connected to the band does not automatically authorize a custom backend to access that account.

Google currently identifies Google Health API as the successor to Fitbit Web API. Its documentation states that new projects are not being onboarded at this time. Legacy Fitbit API support ended on 30 September 2026 and the documented shutdown is 30 October 2026. Therefore a new implementation should target Google Health API, subject to project eligibility, rather than depend on a short lived legacy integration. [1, 2]

The Google Health application, Google Health cloud API, Google Fit and Android Health Connect are distinct products. This architecture uses the Google Health cloud API. An Android Health Connect integration would need an Android component and a separate upload path; it is not a direct replacement for the Python cloud client.

Feasibility checklist
- Confirm whether the team already has an eligible Google Cloud project with Health API access.
- Configure an OAuth client and approved redirect URI; check consent screen and current access requirements.
- Obtain participant authorization for the required read only data bundles.
- Make a real request for each required metric and inspect its structure, units, timestamps and source identifiers.
- Check that sleep details and heart rate resolution are sufficient for the intended screens.
- Confirm renewal behavior and consent duration cover the study, including final reconciliation.

Fallback plan
If access is unavailable, use participant supplied JSON files obtained through an authorized export or an already working integration. Inspect an actual export before committing to its capabilities; exported data may differ from API JSON and may not include every chart field. Build a separate import adapter with file validation, source tagging and duplicate protection. If neither API access nor suitable export is available, postpone real collection and use clearly labeled synthetic fixtures for development.

Implementation decision gate
Proceed with the study only after one real sleep record, one heart rate record and a step record can be imported and read back from the database. Do not assume OAuth setup alone resolves Health API eligibility. Recheck official availability immediately before development.'''),
('System architecture and data flow', '''The frontend is responsible for interaction and visualization. FastAPI owns authentication, validation, synchronization control and summary endpoints. PostgreSQL stores measurements and provenance. A separate scheduled process retrieves provider data; the browser never receives Google access or refresh tokens.

Logical data flow
Fitbit Inspire 3 -> phone synchronization -> Google Health cloud -> authorized Python ingestion -> PostgreSQL -> FastAPI -> React dashboard.

The band and phone handle device synchronization. A successful backend request only retrieves records already present in the cloud; it does not force the band to upload. The interface must distinguish cloud retrieval time from measurement time and device synchronization when that information is available.

Backend modules
- Provider adapter: OAuth supported requests, pagination, time filters and provider response handling.
- Import service: validate JSON files and translate their schema into the same internal model.
- Normalization service: convert units and timestamps, preserve source identity and flag questionable records.
- Repository layer: database transactions, queries and upserts.
- Analytics service: daily summaries, study week summaries and comparison calculations.
- API layer: authenticated routes and a stable response contract for React.

React should fetch dashboard endpoints from FastAPI after sign in. FastAPI queries stored records and returns summary data and chart points. A manual refresh starts a bounded synchronization job and returns its identifier; the frontend checks the status and then refetches stored summaries.

For this scale, one backend, one database and one scheduled command are sufficient. Use a database lock or advisory lock to prevent scheduled and manual jobs from running simultaneously for the same participant. Avoid introducing Redis, Celery or a streaming platform unless later requirements justify them.

Use an adapter boundary so API response changes do not change dashboard components. Keep provider JSON and frontend JSON as different contracts. Store calculation version and source provenance with summaries so the final report can reproduce which logic produced a result.'''),
('Tools and technology choices', '''These are recommended tools, not tools already installed or implemented in the repository. Select mutually compatible supported releases when development begins, then lock dependency versions. Do not copy credentials into the repository.

Frontend tools
- React: reusable components for cards, filters, tables and views. [6]
- TypeScript: typed response models and component props to catch common integration mistakes. [6]
- Vite: development server and production frontend build. [7]
- Recharts: proposed React charts for line, bar and stacked stage views.
- TanStack Query: proposed request caching, loading states and refetch after synchronization.
- CSS Modules: scoped styling without requiring a large design framework.
- Vitest and React Testing Library: component behavior and calculation presentation tests.

Backend and database tools
- Python and FastAPI: HTTP routes, request validation and generated OpenAPI documentation. [8]
- Pydantic: validated settings, request models and response models.
- Uvicorn: ASGI server for FastAPI.
- HTTPX: HTTP client with explicit timeouts for provider requests.
- Google authentication libraries: supported token and OAuth handling, supplemented by the appropriate OAuth flow library.
- SQLAlchemy and psycopg: relational models, transactions and PostgreSQL connectivity.
- Alembic: version controlled database schema migrations.
- PostgreSQL: relational storage plus JSONB for original structured responses. JSONB preserves JSON structure rather than exact whitespace or original byte formatting. [9]
- pytest: backend and ingestion tests using recorded sanitized fixtures.

Development and operation tools
Use Git and GitHub for source history, reviews and issue tracking; Visual Studio Code for editing; FastAPI Swagger UI or Postman for request inspection; and pgAdmin or DBeaver for database inspection. Windows Task Scheduler can run ingestion on a local Windows machine. A hosted environment needs its own scheduler. Docker Compose is optional for reproducible backend and database setup, and requires checking the applicable Docker Desktop license.

Node.js and npm are needed to build React. Python virtual environments isolate backend packages. Use Ruff for Python checks and ESLint for TypeScript checks. Use Playwright for a small number of end to end dashboard journeys. A local study does not require paid hosting or a complex cloud architecture.'''),
('Measurements and interpretation', '''Store measurement units explicitly: heart rate in beats per minute, steps as a count and sleep durations as seconds. Convert sleep durations to hours and minutes only for display. Preserve absent fields as null rather than fabricating values.

Sleep data
Store each sleep session, its start and end, provider identifiers, any supplied asleep duration and available stage intervals. Display main sleep separately from naps. Assign the main sleep session to its local wake date; a session beginning Monday evening and ending Tuesday morning belongs to Tuesday. If the provider does not identify main sleep, choose the longest session ending that date and mark that classification as derived.

The Google sleep guide describes primary stage intervals and separate short awakening intervals that may overlap the stages. Preserve both and use documented semantics when calculating durations; do not add overlapping awakening durations to primary stage totals. [5] If detailed stages are absent, display the available session duration and an explanation.

Sleep quality should be described through duration, continuity, available stages and timing consistency. A proprietary sleep score is optional and requires confirming it is returned for this account. Do not reconstruct or label a custom number as the Fitbit or Google sleep score. A morning self rating from 1 to 5 is an optional separate participant reported measure, with its own consent and input field.

Heart rate data
Store timestamped heart rate observations and the provider supplied daily resting heart rate, when available. Resting heart rate and mean observed heart rate are different measures. Do not replace a missing resting value with the daily average. Chart gaps where samples are absent and do not imply continuous monitoring from sparse data.

Step data
Choose either a Fitbit source filtered stream for a band specific study or the reconciled account stream for values intended to match the application. Record that choice. Never add overlapping device streams together. Sum interval step counts or store a provider daily total; do not sum repeated cumulative totals.

Wearable values are estimates. The dashboard presents recorded patterns without assigning clinical meaning. Include source, coverage and units in cards and tooltips so changes can be interpreted alongside synchronization and wear gaps.'''),
('Database design', '''Use PostgreSQL with normalized tables for querying and JSONB records for provenance. Every measurement belongs to an internal participant identifier and the configured study. Keep Google account identity in the restricted connection record rather than displaying an email on charts.

Core tables
- participants: id UUID primary key, participant_code unique, timezone, created_at.
- studies: id, participant_id foreign key, start_date, end_date_exclusive, timezone, status, calculation_version. Require a 14 day interval.
- connections: id, participant_id, provider_user_id, granted_scopes, encrypted_refresh_token, token_expiry and connection_status. Store encryption keys separately.
- ingestion_runs: id, study_id, started_at, finished_at, mode, status, requested_range, row_count and sanitized error_code.
- raw_records: id, run_id, provider, resource_key, payload_hash, fetched_at, source_updated_at when available, schema_version and payload JSONB. Store only required metric responses.
- sleep_sessions: id, study_id, provider_session_key, source_id, start_at, end_at, wake_date, asleep_seconds nullable, in_bed_seconds nullable, is_main, classification_origin and raw_record_id.
- sleep_segments: id, sleep_session_id, segment_key, start_at, end_at, stage_label and segment_kind for primary stage or short awakening.
- heart_rate_samples: id, study_id, provider_sample_key, source_id, observed_at, bpm and raw_record_id.
- step_intervals: id, study_id, provider_interval_key, source_id, start_at, end_at, step_count and raw_record_id.
- daily_summaries: study_id and local_date composite key, steps nullable, main_sleep_seconds nullable, total_sleep_seconds nullable, resting_bpm nullable, observed_mean_bpm nullable, sample_count, coverage_flags, revision and calculated_at.

Integrity and indexes
Use foreign keys, end greater than start checks and nonnegative duration and step checks. Index heart rate by study_id and observed_at; steps by study_id and start_at; sleep by study_id and wake_date. Add unique constraints on study, provider resource identity and source. If sample IDs are unavailable, define and test a composite source and timestamp key. A timestamp alone may collide across sources.

Upsert revised records instead of appending duplicates. Replace child segments transactionally when a sleep session changes. Keep raw revisions by resource key and payload hash. Derived daily summaries can be rebuilt from stored measurements. Record provider daily totals separately from derived interval totals when both are retained; pick one canonical total and use the other for reconciliation.

At one minute resolution, 14 days contain 20,160 minute slots per metric. This is a sizing illustration, not a promise about actual heart rate sampling. PostgreSQL comfortably supports this prototype scale; storage planning should also include raw revisions and backups.'''),
('OAuth and ingestion workflow', '''Use Google OAuth authorization code flow for server side access, with secure state validation and PKCE where supported. Request the documented minimum read only scopes for sleep, health metrics and activity. Use the official complete scope URIs; scope suffixes in data type tables are not complete authorization strings. Health scopes are categorized as restricted and project review requirements must be checked. [2, 3]

Connection sequence
The participant selects Connect account. FastAPI generates a short lived authorization request tied to the current session. Google handles sign in and consent. The backend validates the callback and exchanges the code, stores tokens securely, records granted permissions and starts the first bounded import. Dashboard sign in and provider data consent are separate concerns even if the same Google identity is used.

Provider requests
The documented service endpoint is https://health.googleapis.com. The v4 resource family includes users.dataTypes.dataPoints.list, reconcile and rollup methods. [4] Implement requests using the relevant metric guides and filters; do not assume every data type supports every operation. Query small date ranges and follow all pagination tokens. Resolve the authorized user resource according to the current API documentation.

Synchronization algorithm
1. Acquire the participant ingestion lock and create an ingestion_runs entry.
2. Renew an expired token using the supported library. If consent is revoked, flag reconnect required.
3. Retrieve required metrics for the study window, respecting provider range limits.
4. Persist raw records; validate and normalize records into database transactions.
5. Upsert measurements using stable keys, handling revised sessions and repeated payloads.
6. Recompute affected daily summaries and record row counts and per metric outcomes.
7. Finish the run, release the lock and expose its status to the dashboard.

For this small dataset, requery the complete study window daily rather than maintain a complicated incremental cursor. While collection is active, do not request future days. This captures delayed uploads and corrections. Run after the participant's morning phone sync and optionally after an evening sync. A computer must be awake and online for a local scheduler to execute.

On HTTP 429 or temporary server errors, use bounded retries with backoff and honor Retry After where provided. A 403 requires checking access and scopes. Do not repeatedly retry invalid requests. Preserve stored data after failures and label the dashboard stale. Commit each coherent metric batch and report partial failures explicitly.'''),
('Daily and weekly calculation rules', '''Configure a study timezone, initially Europe Amsterdam using the IANA identifier Europe/Amsterdam. Store absolute instants as timezone aware PostgreSQL timestamps and preserve provider civil date metadata. Construct each day's midnight boundaries in the study timezone before converting them to UTC. A local day can contain 23 or 25 hours during daylight saving changes.

Period definitions
Day 1 is the configured start date. Week 1 comprises Days 1 through 7 and Week 2 comprises Days 8 through 14. These are study weeks, not necessarily Monday to Sunday calendar weeks. Use half open time intervals, including the start and excluding the end. The end date is start date plus 14 local days.

Daily metrics
Steps: sum canonical, nonoverlapping step intervals belonging to that local day, or use the selected authoritative provider total. Split boundary crossing intervals only using supported finer data or an explicitly documented allocation method.
Main sleep: use the chosen main session ending on that date. Total sleep: add distinct sessions ending on the date, including naps, without overlapping double counts.
Sleep efficiency: 100 times asleep_seconds divided by in_bed_seconds, only when both fields have compatible definitions and the denominator is positive. Otherwise return null.
Observed mean heart rate: arithmetic mean of valid samples, explicitly labeled as a sample mean. If irregular sampling makes a time weighted mean preferable, specify a capped hold interval and test it; never carry a sample across long missing periods.
Resting heart rate: provider daily value, not a derived minimum or sample mean.

Weekly metrics
Show total observed steps, mean steps per valid day, mean main sleep per valid night and mean resting heart rate per valid day. Each metric uses its own valid count. For example, 39,000 observed steps across six valid days gives 6,500 per valid day, not 39,000 divided by seven. Label that week incomplete.

Comparison rules
Week 2 minus Week 1 gives an absolute difference. Percentage change equals 100 times that difference divided by Week 1 only when Week 1 is nonzero and both values are valid. Compare averages when coverage differs; present totals with valid day counts. Missing dates remain present in the 14 day series with null values. True observed zero steps remain zero.

Coverage rules
Track retrieval completeness and observation coverage separately. A complete API response does not prove all day wear. If sampling supports it, count distinct covered minute bins and divide by actual local day minutes. Flag provisional days, missing sessions and incomplete metric retrieval. Do not infer full coverage simply from nonzero steps.'''),
('Dashboard and backend API specification', '''Daily view
Provide a date selector constrained to the study, a last retrieval indicator and cards for main sleep duration, steps, resting heart rate and observed mean heart rate. Below the cards show a heart rate line chart with local time labels and visible gaps, an available sleep stage timeline and an optional hourly step chart. Display main sleep and naps separately. Clicking a chart reveals time, value, unit and source.

Weekly view
Offer Week 1 and Week 2 selectors. Show seven daily bars for steps, seven sleep duration values and a resting heart rate trend. Include averages, valid counts and a completeness label. A comparison panel shows both weeks and absolute changes. The overview shows all 14 dates and study progress.

Interaction and accessibility
Use labeled buttons, readable contrast, keyboard navigation, responsive layouts and a table alternative to charts. Loading, empty, partial, stale and failure states are distinct. Never render missing data as zero height bars without an explanation. A refresh button shows job progress and prevents repeated clicks. A provisional current day should say that values may increase.

Proposed application routes
- GET /api/v1/study: configured study dates, timezone and collection status.
- GET /api/v1/dashboard/daily?date=YYYY-MM-DD: cards, availability flags and main sleep metadata.
- GET /api/v1/dashboard/weekly?week=1: daily series, summaries and valid counts; accept only 1 or 2.
- GET /api/v1/dashboard/comparison: both study weeks and difference calculations.
- GET /api/v1/heart-rate?date=YYYY-MM-DD: chart samples or documented aggregation buckets.
- GET /api/v1/sleep?date=YYYY-MM-DD: sessions and segments for the wake date.
- POST /api/v1/sync: start a bounded authorized ingestion job and return 202 plus run_id.
- GET /api/v1/sync/{run_id}: per metric status and sanitized failures.
- POST /api/v1/import: authenticated JSON import with size and schema limits.
- GET /api/v1/export: participant authorized CSV of daily summaries.

These are project endpoints, not Google endpoints. Include authentication and ownership checks on all data routes. Return 400 for invalid query parameters, 401 for unauthenticated access and 404 for inaccessible resources. Return valid empty responses for dates with no records. Include response schema version, timezone, calculation version and source freshness metadata. Use pagination or server aggregation if chart data is dense.'''),
('Illustrative data contract and repository structure', '''The following response illustrates the project's own daily JSON contract. These values are fictional and do not reproduce a Google response. Raw provider structures must be translated by the adapter.

{
  "schema_version": "1",
  "date": "2026-10-07",
  "timezone": "Europe/Amsterdam",
  "sleep": {"main_asleep_seconds": 25200, "score": null},
  "steps": {"count": 8200, "status": "provisional"},
  "heart_rate": {"resting_bpm": 62, "observed_mean_bpm": 74.2},
  "coverage": {"heart_rate_sample_count": 1100},
  "freshness": {"last_successful_fetch": "2026-10-07T18:00:00Z"},
  "calculation_version": "1"
}

Do not equate sample_count with wear minutes. A separate series endpoint should contain timestamps and bpm values, with null gaps or segments according to the agreed chart contract. Document every status value in OpenAPI and TypeScript models.

Recommended repository layout
frontend/src/components: reusable cards, chart wrappers, status displays and filters.
frontend/src/pages: daily, weekly, comparison and account views.
frontend/src/api: typed API client and query hooks.
frontend/src/types: frontend response models.
backend/app/api: FastAPI routes and authentication dependencies.
backend/app/providers: Google Health and JSON import adapters.
backend/app/services: synchronization, normalization and analytics.
backend/app/models: SQLAlchemy tables.
backend/app/schemas: Pydantic contracts.
backend/app/db: sessions and database configuration.
backend/app/jobs: scheduled ingestion command.
backend/alembic: migrations; backend/tests: fixtures and test suites.
docs: architecture, data dictionary and study protocol.

Configuration
Maintain a committed .env.example with placeholder values for DATABASE_URL, GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET, GOOGLE_REDIRECT_URI, STUDY_START_DATE, STUDY_TIMEZONE, FRONTEND_ORIGIN and collection mode. The actual .env and token encryption key belong outside source control. Store study dates in the database once configured so every service uses the same window.

Use dependency lockfiles and document startup, migrations, account connection, scheduled import, backup and restore commands. A sample dataset should be synthetic and visibly labeled. Keep real exports, database files, tokens and participant reports out of Git.'''),
('Privacy security and deployment', '''Treat sleep, heart rate and activity records as sensitive participant information. Obtain informed participation consent covering what is collected, who sees it, how long it is kept and how withdrawal is handled. Agree the retention date and research or coursework requirements with the institution before collection. This is an engineering checklist, not a legal determination.

Minimize data collection to the configured 14 days and necessary metrics. Use a participant code in screenshots and reports. If provider responses contain extra personal fields, exclude them before persistence where possible. Decide whether retaining full raw metric payloads is necessary and restrict access to them.

Secure implementation
- Keep OAuth tokens and client secrets on the backend and encrypt stored refresh tokens with a separately held key.
- Use authenticated sessions, secure HttpOnly cookies and CSRF protection for cookie authenticated write operations.
- Validate OAuth state and redirect destinations; prevent arbitrary redirect URLs.
- Authorize study ownership on every endpoint and job status lookup.
- Restrict CORS to the actual frontend origin during development; CORS is not authentication.
- Use TLS for remote access, parameterized database operations and bounded JSON import sizes.
- Redact tokens, personal payloads and account identifiers from logs and error messages.
- Protect the database with a dedicated least privilege account and encrypted backups.

Local deployment
React runs through Vite for development and FastAPI runs through Uvicorn. PostgreSQL runs locally or in an optional container. Bind development services to localhost unless network access is intentionally configured. The collection machine must remain available for scheduled jobs; missed runs should be recovered through the next full window query.

Hosted demonstration
Serve the built frontend and proxy /api to FastAPI behind HTTPS. Deploy PostgreSQL privately and configure a separate scheduler. Use environment managed secrets, database migrations, backups and health checks. Check region, service terms and institutional requirements before uploading participant data. Public hosting is optional, and a demonstration can use synthetic data.

Completion and deletion
End measurement collection at the study boundary. Allow a short, agreed reconciliation period to retrieve late records whose measurements fall inside the original window. Then freeze summaries, stop scheduling and revoke unused provider authorization. Retaining existing records after collection is a separate decision. Delete records and expire backups according to the agreed retention policy.'''),
('Two week collection and implementation plan', '''Preparation takes place before the observation period. Development time and recording time are different: do not spend the first study days discovering whether access works.

Before Day 1
Confirm access or a viable export path; obtain consent; define the start date, timezone and exact end boundary. Synchronize the band and verify battery and wear practices. Test actual JSON ingestion, storage and a displayed daily summary. Set up the scheduler and backup. Explain that gaps and charging periods should be recorded in a simple study log.

Days 1 through 7
The participant wears the band during ordinary activity and sleep, following device guidance, and synchronizes the phone each morning and evening. The ingestion job retrieves eligible data. The team checks per metric completion, compares a small number of values with the application and resolves import problems without editing measurements to fit expectations. At the end of Day 7, prepare the Week 1 summary with valid counts.

Days 8 through 14
Continue the same collection routine and ingestion settings. Monitor delayed uploads and revised sleep records. Keep source selection and calculation rules consistent. At the end of Day 14, stop extending the measurement window; retrieve late uploads for in window records during the agreed reconciliation period. Freeze the final dataset and produce Week 2 and comparison summaries.

Proposed development phases
Phase 1: feasibility proof and one real sample per metric.
Phase 2: project scaffolding, database migrations and provider or import adapter.
Phase 3: repeatable ingestion, daily calculations and validated API contracts.
Phase 4: daily and weekly React views, missing data states and accessibility.
Phase 5: integration tests, collection rehearsal and deployment setup.
Phase 6: supervised 14 day collection, final reconciliation and reporting.

Responsibilities
The participant handles wearing and phone synchronization. A backend owner manages provider access, data models and ingestion. A frontend owner implements presentation and interactions. A study owner manages consent, collection checks and reporting. In a small team one person may hold several roles, but each responsibility still needs an owner.

Costs and resources
Budget for the existing band, a compatible phone, internet and an available development machine. The proposed open source application stack can run locally without a hosting subscription. Cloud hosting, domain names, managed databases, institutional reviews and optional product subscriptions may introduce costs; verify prices only if selecting them. Do not assume a paid subscription grants developer API access.'''),
('Validation risks and acceptance criteria', '''Testing must prove the data pipeline and calculations rather than only check that charts appear. Use sanitized recorded provider fixtures and synthetic edge cases. Do not call real health services from ordinary automated tests.

Required checks
- Ingesting the same payload twice produces the same measurement counts and totals.
- A revised sleep record replaces its stages and updates the affected daily summary.
- Pagination retrieves every page; partial failure marks the run incomplete.
- Overnight sleep is assigned to the wake date, with naps handled consistently.
- Short awakenings and overlapping sources do not inflate totals.
- Local midnight and daylight saving boundaries produce correct study day membership.
- Missing metrics are null; verified zero steps remain zero; invalid records are quarantined.
- Weekly averages use per metric valid counts and zero baselines suppress percentage change.
- Expired credentials renew correctly; revoked access requires reconnection; unauthorized API requests cannot read participant data.
- The daily selector, week selector, empty states, chart table alternative and export work end to end.
- Database backup restoration reproduces the same summaries.

Primary risks and responses
API ineligibility: complete feasibility first and maintain the authorized file import path.
Late cloud synchronization: show freshness and repeat the bounded study query.
Incomplete wear: preserve gaps and annotate coverage rather than invent readings.
Provider schema changes: isolate mapping and test recorded response fixtures.
Mixed data sources: select and record a canonical stream and deduplicate by source identity.
Incorrect weekly comparison: show valid counts and compare like denominators.
Privacy exposure: restrict data access, sanitize logs and use synthetic public demonstrations.

Acceptance criteria
The dashboard can display all 14 configured dates, including dates with no data. Every available metric can be traced to a stored source record. Repeated synchronization cannot inflate totals. The two weeks contain exactly seven study dates each. The system operates from stored data during provider outages. The interface labels provisional, missing and stale records. All core calculations and timezone rules pass tests. Setup and recovery instructions allow another team member to run the project. The final report distinguishes observed patterns, missing data and limitations.

Final analysis
Report valid nights, valid step days and valid heart rate days separately. Describe changes between the weeks and explain whether gaps affect interpretation. Optional scatter plots of prior day steps against next morning sleep can illustrate an association, but at most 14 observations and confounding factors make causal or general conclusions inappropriate.'''),
('Official references and implementation decisions', '''References checked on 7 October 2026. Recheck provider availability, scopes and response schemas before implementation. The recommended application architecture and study rules are project design choices.

[1] Google Health API overview and availability
https://developers.google.com/health

[2] About Google Health API and migration timeline
https://developers.google.com/health/about

[3] Google Health API data types and read scopes
https://developers.google.com/health/data-types

[4] Google Health API REST reference and list method
https://developers.google.com/health/reference/rest
https://developers.google.com/health/reference/rest/v4/users.dataTypes.dataPoints/list

[5] Google Health API sleep sessions and awakenings
https://developers.google.com/health/data-types/sleep

[6] React quick start and TypeScript
https://react.dev/learn
https://react.dev/learn/typescript

[7] Vite getting started
https://vite.dev/guide/

[8] FastAPI features and SQL database integration
https://fastapi.tiangolo.com/features/
https://fastapi.tiangolo.com/tutorial/sql-databases/

[9] PostgreSQL JSON types
https://www.postgresql.org/docs/current/datatype-json.html

Decisions to record before development
Confirm the actual Google Health project access and available JSON samples; select Fitbit only or reconciled account data; configure the start date; agree the main sleep and nap rule; define required versus optional fields; choose local or hosted deployment; assign owners; and approve consent and retention arrangements through the institution's normal process.

Recommended first milestone
Demonstrate one authorized JSON retrieval or import, one database write and one React card reading through FastAPI. Once that path works for sleep, heart rate and steps, expand to repeatable ingestion and the daily and weekly dashboard.''')
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.07
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.core_properties.title = 'Sleep Quality Heart Rate and Steps Dashboard Project Specification'
doc.core_properties.subject = 'React Python and PostgreSQL dashboard for a two week Fitbit study'
doc.add_paragraph('Sleep Quality Heart Rate and Steps Dashboard', 'Title')
doc.add_paragraph('Project description and technical implementation specification', 'Subtitle')
doc.add_paragraph('Two week individual study using Fitbit Inspire 3\nPrepared 7 October 2026')
doc.add_paragraph('This report defines the project, recommended tools, data pipeline, database, dashboard behavior and collection protocol. The first milestone is to confirm authorized access to real JSON data. After that, the team can implement a database backed dashboard with reproducible daily and weekly summaries.')
doc.add_heading('Document contents', 1)
for i,(title,_) in enumerate(sections,1):
    doc.add_paragraph(f'{i}  {title}')
for i,(title,body) in enumerate(sections,1):
    doc.add_page_break()
    doc.add_heading(f'{i}  {title}', 1)
    for block in body.split('\n\n'):
        if block.startswith('{'):
            p=doc.add_paragraph(block)
            for r in p.runs: r.font.name='Consolas'; r.font.size=Pt(9)
        elif block.startswith('- '):
            for line in block.splitlines(): doc.add_paragraph(line[2:], 'List Bullet')
        elif '\n- ' in block:
            lead,*lines=block.splitlines()
            p=doc.add_paragraph(); p.add_run(lead).bold=True
            for line in lines: doc.add_paragraph(line[2:] if line.startswith('- ') else line, 'List Bullet')
        else:
            doc.add_paragraph(block)
footer=sec.footer.paragraphs[0]
footer.alignment=2
run=footer.add_run()
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); run._r.addnext(field)
doc.save(ROOT/'Sleep_Dashboard_Project_Specification.docx')
md='# Sleep Quality Heart Rate and Steps Dashboard\n\nProject specification | 7 October 2026\n\n'
md+='\n\n'.join(f'## {i} {t}\n\n{b}' for i,(t,b) in enumerate(sections,1))
(ROOT/'Sleep_Dashboard_Project_Specification.md').write_text(md,encoding='utf-8')
print('Created DOCX and Markdown project specification')
